"""Table detection and export module."""
import csv
import logging
from typing import List, Optional

import cv2
import numpy as np

logger = logging.getLogger(__name__)

try:
    from docx import Document as DocxDocument
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False
    logger.warning("python-docx not available; TableExporter.to_docx_table disabled")


class TableDetector:
    """Detects and reconstructs tables in document images."""

    def __init__(self, config: Optional[dict] = None) -> None:
        """Initialise the table detector.

        Args:
            config: Optional configuration dictionary.
        """
        self._config = config or {}
        logger.debug("TableDetector initialised")

    def detect_grid_lines(self, image: np.ndarray) -> dict:
        """Detect horizontal and vertical grid lines in the image.

        Args:
            image: BGR or grayscale numpy array.

        Returns:
            Dict with keys ``horizontal`` and ``vertical``, each a list of
            y- or x-coordinate integers respectively.
        """
        try:
            gray = (
                cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
                if len(image.shape) == 3
                else image.copy()
            )
            # Binarise
            _, binary = cv2.threshold(
                gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU
            )
            h, w = binary.shape

            # Horizontal lines
            h_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (w // 10, 1))
            h_lines_img = cv2.morphologyEx(binary, cv2.MORPH_OPEN, h_kernel)
            h_coords = self._extract_line_coords(h_lines_img, axis=0)

            # Vertical lines
            v_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, h // 10))
            v_lines_img = cv2.morphologyEx(binary, cv2.MORPH_OPEN, v_kernel)
            v_coords = self._extract_line_coords(v_lines_img, axis=1)

            logger.debug(
                f"Grid lines detected: {len(h_coords)} horizontal, "
                f"{len(v_coords)} vertical"
            )
            return {"horizontal": h_coords, "vertical": v_coords}
        except Exception as exc:
            logger.error(f"detect_grid_lines failed: {exc}")
            return {"horizontal": [], "vertical": []}

    def extract_table_region(
        self, image: np.ndarray, bbox: tuple
    ) -> np.ndarray:
        """Crop a table region from the image.

        Args:
            image: BGR or grayscale numpy array.
            bbox: ``(x1, y1, x2, y2)`` crop coordinates.

        Returns:
            Cropped numpy array.
        """
        try:
            x1, y1, x2, y2 = bbox
            h, w = image.shape[:2]
            x1, y1 = max(0, x1), max(0, y1)
            x2, y2 = min(w, x2), min(h, y2)
            return image[y1:y2, x1:x2]
        except Exception as exc:
            logger.error(f"extract_table_region failed: {exc}")
            return image

    def reconstruct_table(self, image: np.ndarray) -> List[List[str]]:
        """Reconstruct table structure from an image region.

        Detects grid lines, finds cells, and returns a 2-D list of strings.
        OCR integration is handled externally by the pipeline.

        Args:
            image: BGR or grayscale numpy array of the table region.

        Returns:
            2-D list of strings (empty strings as placeholders).
        """
        try:
            grid = self.detect_grid_lines(image)
            h_lines = grid["horizontal"]
            v_lines = grid["vertical"]
            cells = self._find_cells(h_lines, v_lines, image.shape)
            if not cells:
                return [[""]]
            # Determine grid dimensions
            rows = len(h_lines) - 1 if len(h_lines) > 1 else 1
            cols = len(v_lines) - 1 if len(v_lines) > 1 else 1
            table: List[List[str]] = [
                ["" for _ in range(cols)] for _ in range(rows)
            ]
            logger.debug(
                f"Reconstructed table skeleton: {rows}x{cols}"
            )
            return table
        except Exception as exc:
            logger.error(f"reconstruct_table failed: {exc}")
            return [[""]]

    def _find_cells(
        self,
        h_lines: List[int],
        v_lines: List[int],
        image_shape: tuple,
    ) -> List[tuple]:
        """Compute cell bounding boxes from grid line intersections.

        Args:
            h_lines: Sorted y-coordinates of horizontal lines.
            v_lines: Sorted x-coordinates of vertical lines.
            image_shape: ``(height, width[, channels])`` tuple.

        Returns:
            List of ``(x1, y1, x2, y2)`` tuples.
        """
        cells: List[tuple] = []
        try:
            h = h_lines if h_lines else [0, image_shape[0]]
            v = v_lines if v_lines else [0, image_shape[1]]
            for row_idx in range(len(h) - 1):
                for col_idx in range(len(v) - 1):
                    cells.append((
                        v[col_idx], h[row_idx],
                        v[col_idx + 1], h[row_idx + 1],
                    ))
        except Exception as exc:
            logger.error(f"_find_cells failed: {exc}")
        return cells

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_line_coords(mask: np.ndarray, axis: int) -> List[int]:
        """Extract centre coordinates of lines from a binary mask.

        Args:
            mask: Binary image where lines are white.
            axis: 0 for horizontal lines (y-coords), 1 for vertical (x-coords).

        Returns:
            Sorted list of integer coordinates.
        """
        projection = mask.sum(axis=1 - axis)
        threshold = projection.max() * 0.5 if projection.max() > 0 else 1
        coords: List[int] = []
        in_line = False
        start = 0
        for i, val in enumerate(projection):
            if val >= threshold and not in_line:
                in_line = True
                start = i
            elif val < threshold and in_line:
                in_line = False
                coords.append((start + i) // 2)
        return sorted(coords)


class TableExporter:
    """Exports reconstructed tables to various formats."""

    def to_csv(self, table: List[List[str]], path: str) -> None:
        """Write a 2-D table to a CSV file.

        Args:
            table: 2-D list of strings.
            path: Destination file path.
        """
        try:
            import pathlib
            pathlib.Path(path).parent.mkdir(parents=True, exist_ok=True)
            with open(path, "w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(fh)
                writer.writerows(table)
            logger.info(f"Table exported to CSV: {path}")
        except Exception as exc:
            logger.error(f"to_csv failed: {exc}")

    def to_docx_table(self, table: List[List[str]], doc) -> None:
        """Add a table to a python-docx Document object.

        Args:
            table: 2-D list of strings.
            doc: A :class:`docx.Document` instance.
        """
        if not _DOCX_AVAILABLE:
            logger.warning("python-docx not available; cannot add table")
            return
        try:
            rows = len(table)
            cols = max((len(row) for row in table), default=1)
            tbl = doc.add_table(rows=rows, cols=cols)
            tbl.style = "Table Grid"
            for r_idx, row in enumerate(table):
                for c_idx, cell_text in enumerate(row):
                    tbl.cell(r_idx, c_idx).text = str(cell_text)
            logger.debug("Table added to DOCX document")
        except Exception as exc:
            logger.error(f"to_docx_table failed: {exc}")
